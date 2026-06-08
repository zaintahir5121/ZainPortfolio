from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

ACCENT = RGBColor(0x1e, 0x3a, 0x5f)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF5, 0xF5, 0xF5)
DARK_TEXT = RGBColor(0x1A, 0x1A, 0x2E)
MID_GRAY = RGBColor(0x55, 0x55, 0x55)


def set_cell_bg(cell, color_hex):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        tag = OxmlElement(f'w:{edge}')
        tag.set(qn('w:val'), kwargs.get(edge, 'none'))
        if kwargs.get(edge, 'none') != 'none':
            tag.set(qn('w:sz'), kwargs.get('sz', '4'))
            tag.set(qn('w:color'), kwargs.get('color', '1e3a5f'))
        tcBorders.append(tag)
    tcPr.append(tcBorders)


def add_horizontal_rule(doc, color='1e3a5f', thickness=12):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(2)
    para.paragraph_format.space_after = Pt(4)
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(thickness))
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)
    return para


def add_section_heading(doc, text):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(12)
    para.paragraph_format.space_after = Pt(0)
    run = para.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = ACCENT
    run.font.name = 'Calibri'
    add_horizontal_rule(doc)
    return para


def add_bullet(doc, text, bold_part=None):
    para = doc.add_paragraph(style='List Bullet')
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(1)
    para.paragraph_format.left_indent = Inches(0.25)
    if bold_part and text.startswith(bold_part):
        run = para.add_run(bold_part)
        run.bold = True
        run.font.size = Pt(10)
        run.font.name = 'Calibri'
        run.font.color.rgb = DARK_TEXT
        rest = text[len(bold_part):]
        run2 = para.add_run(rest)
        run2.font.size = Pt(10)
        run2.font.name = 'Calibri'
        run2.font.color.rgb = DARK_TEXT
    else:
        run = para.add_run(text)
        run.font.size = Pt(10)
        run.font.name = 'Calibri'
        run.font.color.rgb = DARK_TEXT
    return para


def set_para_font(para, size=10, color=None, bold=False, name='Calibri'):
    for run in para.runs:
        run.font.size = Pt(size)
        run.font.name = name
        run.bold = bold
        if color:
            run.font.color.rgb = color


def add_job(doc, title, company, period, location, bullets, tech):
    # Title row
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(8)
    para.paragraph_format.space_after = Pt(1)
    r1 = para.add_run(title)
    r1.bold = True
    r1.font.size = Pt(11)
    r1.font.color.rgb = DARK_TEXT
    r1.font.name = 'Calibri'
    r2 = para.add_run(f'  |  {company}')
    r2.bold = True
    r2.font.size = Pt(11)
    r2.font.color.rgb = ACCENT
    r2.font.name = 'Calibri'

    # Period + location
    meta = doc.add_paragraph()
    meta.paragraph_format.space_before = Pt(0)
    meta.paragraph_format.space_after = Pt(3)
    rm = meta.add_run(f'{period}   •   {location}')
    rm.font.size = Pt(9)
    rm.font.color.rgb = MID_GRAY
    rm.font.name = 'Calibri'
    rm.italic = True

    for b in bullets:
        add_bullet(doc, b)

    # Tech line
    tech_para = doc.add_paragraph()
    tech_para.paragraph_format.space_before = Pt(2)
    tech_para.paragraph_format.space_after = Pt(2)
    tech_para.paragraph_format.left_indent = Inches(0.25)
    rt1 = tech_para.add_run('Tech: ')
    rt1.bold = True
    rt1.font.size = Pt(9)
    rt1.font.color.rgb = ACCENT
    rt1.font.name = 'Calibri'
    rt2 = tech_para.add_run(tech)
    rt2.font.size = Pt(9)
    rt2.font.color.rgb = MID_GRAY
    rt2.font.name = 'Calibri'


doc = Document()

# Page margins
for section in doc.sections:
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

# ── HEADER TABLE (name + contact) ──────────────────────────────────────────
header_table = doc.add_table(rows=1, cols=2)
header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
header_table.style = 'Table Grid'

left_cell = header_table.cell(0, 0)
right_cell = header_table.cell(0, 1)

# Set background
set_cell_bg(left_cell, '1e3a5f')
set_cell_bg(right_cell, '1e3a5f')

# Remove borders
for cell in [left_cell, right_cell]:
    set_cell_border(cell)

# Set column widths
header_table.columns[0].width = Inches(3.8)
header_table.columns[1].width = Inches(3.0)

# Left: Name & Title
lp1 = left_cell.paragraphs[0]
lp1.paragraph_format.space_before = Pt(10)
lp1.paragraph_format.space_after = Pt(2)
lp1.paragraph_format.left_indent = Pt(8)
r_name = lp1.add_run('Zain Abbas Tahir')
r_name.bold = True
r_name.font.size = Pt(22)
r_name.font.color.rgb = WHITE
r_name.font.name = 'Calibri'

lp2 = left_cell.add_paragraph()
lp2.paragraph_format.space_before = Pt(0)
lp2.paragraph_format.space_after = Pt(10)
lp2.paragraph_format.left_indent = Pt(8)
r_title = lp2.add_run('AI Technical Lead & Cloud Architect')
r_title.font.size = Pt(12)
r_title.font.color.rgb = RGBColor(0xA8, 0xC8, 0xE8)
r_title.font.name = 'Calibri'

# Right: Contact info
contact_items = [
    ('Location:', 'Kuala Lumpur, Malaysia'),
    ('Email:', 'zabbastahir@gmail.com'),
    ('LinkedIn:', 'linkedin.com/in/zainabbastahir'),
    ('GitHub:', 'github.com/zainabbastahir'),
    ('Upwork:', 'upwork.com/freelancers/~013b69c81fcb1ae708'),
]

first_cp = right_cell.paragraphs[0]
first_cp.paragraph_format.space_before = Pt(8)
first_cp.paragraph_format.space_after = Pt(1)
first_cp.paragraph_format.left_indent = Pt(4)
r_label = first_cp.add_run(contact_items[0][0] + ' ')
r_label.bold = True
r_label.font.size = Pt(9)
r_label.font.color.rgb = RGBColor(0xA8, 0xC8, 0xE8)
r_label.font.name = 'Calibri'
r_val = first_cp.add_run(contact_items[0][1])
r_val.font.size = Pt(9)
r_val.font.color.rgb = WHITE
r_val.font.name = 'Calibri'

for label, val in contact_items[1:]:
    cp = right_cell.add_paragraph()
    cp.paragraph_format.space_before = Pt(1)
    cp.paragraph_format.space_after = Pt(1)
    cp.paragraph_format.left_indent = Pt(4)
    rl = cp.add_run(label + ' ')
    rl.bold = True
    rl.font.size = Pt(9)
    rl.font.color.rgb = RGBColor(0xA8, 0xC8, 0xE8)
    rl.font.name = 'Calibri'
    rv = cp.add_run(val)
    rv.font.size = Pt(9)
    rv.font.color.rgb = WHITE
    rv.font.name = 'Calibri'

# padding bottom
cp_last = right_cell.add_paragraph()
cp_last.paragraph_format.space_before = Pt(2)
cp_last.paragraph_format.space_after = Pt(8)

# ── PROFESSIONAL SUMMARY ───────────────────────────────────────────────────
add_section_heading(doc, 'Professional Summary')
summary_para = doc.add_paragraph()
summary_para.paragraph_format.space_before = Pt(4)
summary_para.paragraph_format.space_after = Pt(6)
rs = summary_para.add_run(
    '13+ years of experience building intelligent, scalable, and predictive systems with RAG, LLMs, Azure Cloud, '
    '.NET Core, and Angular. Proven track record across enterprise, healthcare, fintech, and logistics domains. '
    'Available for remote, freelance, and technical leadership roles.'
)
rs.font.size = Pt(10)
rs.font.name = 'Calibri'
rs.font.color.rgb = DARK_TEXT

# ── KEY STATS ──────────────────────────────────────────────────────────────
stats_table = doc.add_table(rows=1, cols=3)
stats_table.alignment = WD_TABLE_ALIGNMENT.CENTER
stats_table.style = 'Table Grid'

stats = [
    ('13+', 'Years Experience'),
    ('50+', 'Projects Completed'),
    ('20+', 'Enterprise Clients'),
]

for i, (num, label) in enumerate(stats):
    cell = stats_table.cell(0, i)
    set_cell_bg(cell, 'EAF0F8')
    set_cell_border(cell)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    p1 = cell.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(6)
    p1.paragraph_format.space_after = Pt(0)
    rn = p1.add_run(num)
    rn.bold = True
    rn.font.size = Pt(20)
    rn.font.color.rgb = ACCENT
    rn.font.name = 'Calibri'

    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(6)
    rl2 = p2.add_run(label)
    rl2.font.size = Pt(9)
    rl2.font.color.rgb = MID_GRAY
    rl2.font.name = 'Calibri'

# ── EXPERIENCE ─────────────────────────────────────────────────────────────
add_section_heading(doc, 'Professional Experience')

add_job(doc,
    'Technical Lead', 'Aventra Group',
    'Nov 2025 – Present', 'Kuala Lumpur, Malaysia',
    [
        'Leading project teams and defining technical solutions for AI-powered products',
        'Building RAG systems, AI Search, and Predictive AI solutions',
        'Designing microservices architecture and managing Azure cloud infrastructure',
    ],
    'RAG, LLMs, AI Search, .NET Core, Angular, Azure, Microservices, Copilot Studio'
)

add_job(doc,
    'Technical Lead', 'DHL IT Services',
    'Sep 2022 – Oct 2025', 'Kuala Lumpur, Malaysia (Hybrid)',
    [
        'Managed developer teams and Azure cloud infrastructure',
        'Defined and delivered AI-based Route Optimization System',
        'Built RAG System for Global Search across enterprise data',
    ],
    '.NET Core, Angular 18, Azure, RAG, Cosmos DB, Redis, Microservices'
)

add_job(doc,
    'Senior Software Engineer', 'Tapcheck',
    '2019 – 2022', 'Remote (US-based)',
    [
        'Led payroll integrations with major HR platforms',
        'Established coding standards and conducted code reviews',
    ],
    '.NET Core, Angular, REST APIs, SQL Server'
)

add_job(doc,
    'Senior Software Engineer', 'UMCH',
    '2018 – 2019', 'Kuala Lumpur, Malaysia',
    [
        'Built facial and emotion recognition APIs',
        'Developed intelligent chatbots with NLP and sentiment analysis',
    ],
    'Python, Machine Learning, AI/ML, NLP'
)

add_job(doc,
    'Software Architect', 'MTBC (CareCloud)',
    '2016 – 2018', 'Islamabad, Pakistan',
    [
        'Managed complete software lifecycle for healthcare solutions',
        'Mentored junior developers, ensured HIPAA compliance',
    ],
    'ASP.NET MVC, Entity Framework, Healthcare IT, HIPAA'
)

add_job(doc,
    'Software Engineer', 'Interactive Group & Moftak Solutions',
    '2013 – 2016', 'Islamabad, Pakistan',
    [
        'Developed custom web applications and robust APIs',
    ],
    'ASP.NET, C#, JavaScript, SQL Server'
)

# ── TECHNICAL SKILLS ───────────────────────────────────────────────────────
add_section_heading(doc, 'Technical Skills')

skills = [
    ('AI & Machine Learning', 'RAG, LLMs, Predictive AI, Sentiment Analysis, Azure AI Foundry, Copilot Studio'),
    ('Cloud & Infrastructure', 'Microsoft Azure, Azure Service Bus, Data Factory, Kubernetes, Key Vault, AD & RBAC'),
    ('Backend Development', '.NET / .NET Core, Microservices, CQRS, Event-Driven Architecture'),
    ('Frontend Development', 'Angular 17+, Blazor, React'),
    ('DevOps & Automation', 'CI/CD Pipelines, Azure DevOps, Terraform, Docker, TDD, ARM/Bicep'),
    ('Databases', 'SQL Server, EF Core, Cosmos DB, Redis Cache'),
]

skills_table = doc.add_table(rows=len(skills), cols=2)
skills_table.alignment = WD_TABLE_ALIGNMENT.LEFT
skills_table.style = 'Table Grid'

for i, (cat, vals) in enumerate(skills):
    c1 = skills_table.cell(i, 0)
    c2 = skills_table.cell(i, 1)
    bg = 'EAF0F8' if i % 2 == 0 else 'FFFFFF'
    set_cell_bg(c1, '1e3a5f')
    set_cell_bg(c2, bg)
    set_cell_border(c1)
    set_cell_border(c2)

    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_before = Pt(4)
    p1.paragraph_format.space_after = Pt(4)
    p1.paragraph_format.left_indent = Pt(6)
    rc = p1.add_run(cat)
    rc.bold = True
    rc.font.size = Pt(9)
    rc.font.color.rgb = WHITE
    rc.font.name = 'Calibri'

    p2 = c2.paragraphs[0]
    p2.paragraph_format.space_before = Pt(4)
    p2.paragraph_format.space_after = Pt(4)
    p2.paragraph_format.left_indent = Pt(6)
    rv2 = p2.add_run(vals)
    rv2.font.size = Pt(9)
    rv2.font.color.rgb = DARK_TEXT
    rv2.font.name = 'Calibri'

# Set column widths for skills table
for row in skills_table.rows:
    row.cells[0].width = Inches(2.0)
    row.cells[1].width = Inches(4.8)

# ── NOTABLE PROJECTS ───────────────────────────────────────────────────────
add_section_heading(doc, 'Notable Projects')

projects = [
    ('Document Management System',
     'Enterprise DMS with role-based access, multi-tenant, PDF viewing, WhatsApp/email sharing',
     '.NET Core, Azure, Angular'),
    ('AI Customer Support Bot',
     'Context-aware chatbot with RAG and Azure OpenAI',
     'RAG, Azure OpenAI, LLMs'),
    ('Predictive Analytics Engine',
     'Sales forecasting and demand prediction',
     'Azure ML, Data Factory, Power BI'),
    ('Face Recognition Login',
     'Passwordless auth with facial recognition',
     'Python, AI/ML, .NET'),
]

proj_table = doc.add_table(rows=1, cols=2)
proj_table.alignment = WD_TABLE_ALIGNMENT.LEFT
proj_table.style = 'Table Grid'

# We'll do 2 projects per row
for row_idx in range(0, len(projects), 2):
    if row_idx == 0:
        row = proj_table.rows[0]
    else:
        row = proj_table.add_row()

    for col_idx in range(2):
        proj_idx = row_idx + col_idx
        if proj_idx >= len(projects):
            break
        pname, pdesc, ptech = projects[proj_idx]
        cell = row.cells[col_idx]
        set_cell_bg(cell, 'F8FAFD')
        set_cell_border(cell, top='single', left='single', bottom='single', right='single',
                        sz='4', color='1e3a5f')

        pp1 = cell.paragraphs[0]
        pp1.paragraph_format.space_before = Pt(6)
        pp1.paragraph_format.space_after = Pt(2)
        pp1.paragraph_format.left_indent = Pt(6)
        rpp = pp1.add_run(pname)
        rpp.bold = True
        rpp.font.size = Pt(10)
        rpp.font.color.rgb = ACCENT
        rpp.font.name = 'Calibri'

        pp2 = cell.add_paragraph()
        pp2.paragraph_format.space_before = Pt(0)
        pp2.paragraph_format.space_after = Pt(2)
        pp2.paragraph_format.left_indent = Pt(6)
        rpd = pp2.add_run(pdesc)
        rpd.font.size = Pt(9)
        rpd.font.color.rgb = DARK_TEXT
        rpd.font.name = 'Calibri'

        pp3 = cell.add_paragraph()
        pp3.paragraph_format.space_before = Pt(2)
        pp3.paragraph_format.space_after = Pt(6)
        pp3.paragraph_format.left_indent = Pt(6)
        rpt1 = pp3.add_run('Tech: ')
        rpt1.bold = True
        rpt1.font.size = Pt(9)
        rpt1.font.color.rgb = ACCENT
        rpt1.font.name = 'Calibri'
        rpt2 = pp3.add_run(ptech)
        rpt2.font.size = Pt(9)
        rpt2.font.color.rgb = MID_GRAY
        rpt2.font.name = 'Calibri'

for row in proj_table.rows:
    for cell in row.cells:
        cell.width = Inches(3.4)

# ── EDUCATION (brief footer note) ─────────────────────────────────────────
sp = doc.add_paragraph()
sp.paragraph_format.space_before = Pt(12)
sp.paragraph_format.space_after = Pt(2)
re1 = sp.add_run('Education: ')
re1.bold = True
re1.font.size = Pt(10)
re1.font.color.rgb = ACCENT
re1.font.name = 'Calibri'
re2 = sp.add_run('Bachelor of Science in Computer Science  |  COMSATS University Islamabad  |  2009 – 2013')
re2.font.size = Pt(10)
re2.font.color.rgb = DARK_TEXT
re2.font.name = 'Calibri'

add_horizontal_rule(doc, color='AAAAAA', thickness=6)

doc.save('/home/user/ZainPortfolio/Zain_Abbas_Tahir_CV.docx')
print('CV saved successfully.')
